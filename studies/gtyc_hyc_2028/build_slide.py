#!/usr/bin/env python
"""Build an editable PowerPoint slide (native charts and table) for one venue/season window.

    python studies/gtyc_hyc_2028/build_slide.py                       # HYC September
    python studies/gtyc_hyc_2028/build_slide.py --venue gtyc --season sep
    python studies/gtyc_hyc_2028/build_slide.py --venue hyc --season jun --out my.pptx

Everything on the slide is a native PowerPoint object (text boxes, a filled radar chart for the
wind rose, a column chart, a table), so fonts, sizes and colors can be restyled in PowerPoint to
match another deck. The numbers come from the study's outputs, so rerunning after the analysis
keeps the slide in sync. Requires python-pptx (``pip install ".[slides]"``).

PowerPoint has no stacked polar bar chart, so the wind rose is cumulative filled radar series:
each color band shows its own speed bin.
"""
import argparse
import math
import re
from pathlib import Path

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from hrrr_wind.config import load_study
from hrrr_wind.stats import load_series, windrose_table

HERE = Path(__file__).parent
FONT = "Calibri"
INK, MUTED, GRID = "1F2937", "6B7280", "D1D5DB"
BIN_COLORS = ["482677", "3B5B92", "238A8D", "29AF7F", "84D44B"]    # slow to fast
BIN_NAMES = ["0–5 kt", "5–10 kt", "10–15 kt", "15–20 kt", "20+ kt"]
SEASONS = {"sep": ("September", "Sept"), "jun": ("June", "June")}
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August",
          "September", "October", "November", "December"]


def rgb(hex_):
    return RGBColor.from_string(hex_)


def textbox(slide, name, text, x, y, w, h, size, color=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.name, run.font.size, run.font.bold = FONT, Pt(size), bold
    run.font.color.rgb = rgb(color)
    return box


def style_chart_text(chart, size=12):
    chart.font.name = FONT
    chart.font.size = Pt(size)
    chart.font.color.rgb = rgb(INK)


def window_title(study, venue, season, window):
    name = study.venues[venue].name
    club = venue.upper() + (f" ({name[name.index('(') + 1:name.index(')')]})" if "(" in name else "")
    m0, d0 = int(window.core[0][:2]), int(window.core[0][3:])
    m1, d1 = int(window.core[1][:2]), int(window.core[1][3:])
    span = f"{MONTHS[m0 - 1]} {d0}–{d1}" if m0 == m1 else f"{MONTHS[m0 - 1]} {d0} – {MONTHS[m1 - 1]} {d1}"
    return f"{club}: {span} wind, {min(study.years)}–{max(study.years)}"


def build_slide(study, venue, season, out_path):
    window = next(w for w in study.windows if w.venue == venue and w.season == season)
    d = load_series(study, venue, window.cache_tag)
    race = d[(d["md"] >= window.lo) & (d["md"] <= window.hi) & d["hour"].isin(study.race_utc)]
    rose = windrose_table(race["spd_kt"], race["dir"])
    stats = pd.read_csv(study.out_dir / "summary_stats.csv")
    s = stats[(stats.venue == venue) & (stats.season == season) & (stats.scope == "core")].iloc[0]
    v = study.venues[venue]

    values = np.array(rose["values"])
    n_bins = 5 if values[4].sum() > 0 else 4
    cumulative = [values[i:].sum(axis=0) for i in range(n_bins)]
    top = max(2, int(math.ceil(cumulative[0].max() / 2.0) * 2))

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    textbox(slide, "Title", window_title(study, venue, season, window), 0.6, 0.35, 12.1, 0.7, 32, bold=True)
    textbox(slide, "Subtitle",
            f"HRRR 10 m wind, racing hours {re.sub(r'(\d)(AM|PM)', r'\1 \2', study.race_label).replace('-', '–')}, n = {rose['n']} hours. "
            f"Sample point {v.offshore_nm:g} nm out at bearing {v.bearing:g}° from the club",
            0.6, 1.05, 12.1, 0.4, 16, color=MUTED)

    # Wind rose: filled radar with cumulative series (largest first, so each band shows its own bin)
    textbox(slide, "Rose heading", "Wind rose: % of hours by direction wind blows from", 0.6, 1.7, 6.2, 0.35, 16, bold=True)
    cd = CategoryChartData()
    cd.categories = rose["sectors"]
    for i in range(n_bins):
        cd.add_series(BIN_NAMES[i], [round(float(x), 3) for x in cumulative[i]])
    gf = slide.shapes.add_chart(XL_CHART_TYPE.RADAR_FILLED, Inches(0.5), Inches(2.05), Inches(6.4), Inches(5.0), cd)
    gf.name = "Wind rose"
    chart = gf.chart
    chart.has_title = False
    style_chart_text(chart, 12)
    chart.has_legend = True
    chart.legend.position, chart.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
    chart.legend.font.size = Pt(14)
    va = chart.value_axis
    va.minimum_scale, va.maximum_scale, va.major_unit = 0, top, 2
    va.tick_labels.font.size = Pt(11)
    va.tick_labels.font.color.rgb = rgb(MUTED)
    va.major_gridlines.format.line.color.rgb = rgb(GRID)
    chart.category_axis.tick_labels.font.size = Pt(12)
    for ser, color in zip(chart.plots[0].series, BIN_COLORS):
        ser.format.fill.solid()
        ser.format.fill.fore_color.rgb = rgb(color)
        ser.format.line.color.rgb = rgb(color)

    # Speed distribution
    textbox(slide, "Distribution heading", "Speed distribution (% of hours)", 7.3, 1.7, 5.4, 0.35, 16, bold=True)
    cd2 = CategoryChartData()
    cd2.categories = ["0–5", "5–10", "10–15", "15–20", "20+"]
    cd2.add_series("% of hours", [round(float(x), 1) for x in rose["distribution"]])
    gf2 = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7.2), Inches(2.05), Inches(5.6), Inches(2.45), cd2)
    gf2.name = "Speed distribution"
    c2 = gf2.chart
    c2.has_title = False
    style_chart_text(c2, 12)
    c2.has_legend = False
    plot = c2.plots[0]
    plot.gap_width = 40
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = '0"%"', False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size = Pt(12)
    for i, color in enumerate(BIN_COLORS):
        pt = plot.series[0].points[i]
        pt.format.fill.solid()
        pt.format.fill.fore_color.rgb = rgb(color)
    c2.value_axis.visible = False
    c2.value_axis.has_major_gridlines = False
    c2.value_axis.maximum_scale = max(45, math.ceil(max(rose["distribution"]) / 5) * 5 + 5)
    c2.category_axis.has_title = True
    c2.category_axis.axis_title.text_frame.text = "Wind speed (kt)"
    run = c2.category_axis.axis_title.text_frame.paragraphs[0].runs[0]
    run.font.size, run.font.bold, run.font.name = Pt(12), False, FONT
    run.font.color.rgb = rgb(MUTED)

    # Statistics table
    rows = [("Mean / median", f"{s['mean']:.1f} / {s['median']:.1f} kt"),
            ("90th percentile / max", f"{s['p90']:.1f} / {s['max']:.1f} kt"),
            ("Below 5 kt", f"{s['pct_lt5']:.0f}%"),
            ("5–15 kt", f"{s['pct_5_15']:.0f}%"),
            ("15 kt or more", f"{s['pct_15_20'] + s['pct_ge20']:.1f}%"),
            ("Usable days (3+ h in 5–15 kt)", f"{s['usable_day_pct']:.0f}%"),
            ("Direction steadiness R", f"{s['R']:.2f}")]
    shape = slide.shapes.add_table(len(rows), 2, Inches(7.3), Inches(4.65), Inches(5.4), Inches(0.3 * len(rows)))
    shape.name = "Summary statistics"
    table = shape.table
    table.first_row = table.horz_banding = False
    table.columns[0].width, table.columns[1].width = Inches(3.5), Inches(1.9)
    for r, (label, value) in enumerate(rows):
        table.rows[r].height = Inches(0.3)
        for c, (text, bold) in enumerate(((label, False), (value, True))):
            cell = table.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb("F3F4F6" if r % 2 == 0 else "FFFFFF")
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.margin_left = cell.margin_right = Inches(0.1)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.RIGHT if c else PP_ALIGN.LEFT
            run = para.add_run()
            run.text = text
            run.font.name, run.font.size, run.font.bold = FONT, Pt(14), bold
            run.font.color.rgb = rgb(INK)

    textbox(slide, "Source",
            f"Source: NOAA HRRR F00 analysis, water cells within {v.radius_nm:g} nm of the sample point. "
            "Wind rose bands are cumulative filled areas, so each color shows its own speed band.",
            0.6, 7.05, 12.1, 0.3, 11, color=MUTED)
    prs.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[1].strip())
    ap.add_argument("--venue", default="hyc")
    ap.add_argument("--season", default="sep", choices=list(SEASONS))
    ap.add_argument("--out", help="output .pptx (default: slides/<VENUE>_<Season>_wind_slide.pptx)")
    args = ap.parse_args()
    study = load_study(HERE / "study.toml")
    out = Path(args.out) if args.out else HERE / "slides" / f"{args.venue.upper()}_{SEASONS[args.season][1]}_wind_slide.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    print("wrote", build_slide(study, args.venue, args.season, out))


if __name__ == "__main__":
    main()
