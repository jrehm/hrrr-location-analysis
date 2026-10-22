import importlib.util
from pathlib import Path

import pytest

pytest.importorskip("pptx")   # optional dependency (pip install ".[slides]")
from pptx import Presentation  # noqa: E402

from hrrr_wind.config import load_study  # noqa: E402

SCRIPT = Path(__file__).resolve().parents[1] / "studies" / "gtyc_hyc_2028" / "build_slide.py"


def load_builder():
    spec = importlib.util.spec_from_file_location("build_slide", SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def test_slide_is_native_and_matches_study_numbers(tmp_path):
    study = load_study(SCRIPT.parent / "study.toml")
    out = load_builder().build_slide(study, "hyc", "sep", tmp_path / "s.pptx")
    slide = Presentation(out).slides[0]
    shapes = {s.name: s for s in slide.shapes}
    assert shapes["Title"].text_frame.text == "HYC (Caseville): September 8–14 wind, 2015–2025"
    assert "n = 523 hours" in shapes["Subtitle"].text_frame.text
    assert "10 AM–4 PM EDT" in shapes["Subtitle"].text_frame.text
    assert shapes["Wind rose"].has_chart and shapes["Speed distribution"].has_chart
    dist = list(shapes["Speed distribution"].chart.plots[0].series[0].values)
    assert dist[0] == pytest.approx(25.2, abs=0.05) and sum(dist) == pytest.approx(100, abs=0.2)
    cells = [[c.text for c in r.cells] for r in shapes["Summary statistics"].table.rows]
    assert cells[0] == ["Mean / median", "8.7 / 8.1 kt"] and cells[4] == ["15 kt or more", "9.8%"]
    assert not shapes["Wind rose"].chart.has_title and not shapes["Speed distribution"].chart.has_title


def test_other_windows_build(tmp_path):
    study = load_study(SCRIPT.parent / "study.toml")
    build = load_builder().build_slide
    assert Presentation(build(study, "gtyc", "sep", tmp_path / "g.pptx")).slides[0].shapes[0].text_frame.text.startswith("GTYC: September 8–14")
    assert "June 14–21" in Presentation(build(study, "hyc", "jun", tmp_path / "j.pptx")).slides[0].shapes[0].text_frame.text
